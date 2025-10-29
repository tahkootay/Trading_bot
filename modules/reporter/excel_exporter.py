"""
Excel Export Module for Trading Strategies

This module provides Excel export functionality with color-coded visualization
for trading strategy signals, positions, and analysis data.

Features:
- Color-coded position and signal visualization
- Automatic column width adjustment
- Legend worksheet with strategy parameters
- Support for multiple data formats
- Professional Excel formatting
"""

from typing import List, Dict, Optional, Any
import os
from datetime import datetime
import pandas as pd

try:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font
    from openpyxl.utils.dataframe import dataframe_to_rows
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False


class ExcelExporter:
    """
    Professional Excel exporter for trading strategy analysis.
    
    Provides color-coded visualization of trading signals, positions,
    and performance metrics with automatic formatting and legend generation.
    """
    
    def __init__(self):
        """Initialize Excel exporter with default color scheme."""
        self.colors = {
            'LONG': PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid"),    # Light Green
            'BUY': PatternFill(start_color="32CD32", end_color="32CD32", fill_type="solid"),     # Lime Green  
            'SELL': PatternFill(start_color="FF6B6B", end_color="FF6B6B", fill_type="solid"),   # Light Red
            'SHORT': PatternFill(start_color="FFA07A", end_color="FFA07A", fill_type="solid"),  # Light Salmon
            'NONE': PatternFill(start_color="F0F0F0", end_color="F0F0F0", fill_type="solid"),   # Light Gray
            'HOLD': PatternFill(start_color="FFD700", end_color="FFD700", fill_type="solid"),   # Gold
            'HEADER': PatternFill(start_color="4682B4", end_color="4682B4", fill_type="solid")  # Steel Blue
        }
        
        self.header_font = Font(color="FFFFFF", bold=True)
    
    def check_availability(self) -> bool:
        """Check if Excel export is available."""
        return EXCEL_AVAILABLE
    
    def export_strategy_signals(
        self, 
        data: List[Dict[str, Any]], 
        output_path: str,
        strategy_info: Optional[Dict[str, Any]] = None
    ) -> str:
        """
        Export strategy signals to Excel with color coding.
        
        Args:
            data: List of signal data dictionaries
            output_path: Path to save Excel file
            strategy_info: Optional strategy information for legend
        
        Returns:
            Path to saved Excel file
        
        Raises:
            RuntimeError: If openpyxl is not available
            ValueError: If data is empty
        """
        if not EXCEL_AVAILABLE:
            raise RuntimeError(
                "Excel export not available. Install openpyxl: pip install openpyxl"
            )
        
        if not data:
            raise ValueError("No data to export")
        
        # Ensure output directory exists
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        
        # Create DataFrame
        df = pd.DataFrame(data)
        
        # Create workbook and main worksheet
        wb = Workbook()
        ws = wb.active
        ws.title = "Strategy Signals"
        
        # Add headers with formatting
        self._add_headers(ws, df.columns)
        
        # Add data with color coding
        self._add_data_with_colors(ws, df)
        
        # Auto-adjust column widths
        self._adjust_column_widths(ws)
        
        # Add legend worksheet if strategy info provided
        if strategy_info:
            self._add_legend_worksheet(wb, strategy_info)
        
        # Save workbook
        wb.save(output_path)
        
        return output_path
    
    def _add_headers(self, worksheet, columns):
        """Add formatted headers to worksheet."""
        for col_idx, column in enumerate(columns, 1):
            cell = worksheet.cell(row=1, column=col_idx, value=column)
            cell.fill = self.colors['HEADER']
            cell.font = self.header_font
    
    def _add_data_with_colors(self, worksheet, df: pd.DataFrame):
        """Add data rows with color coding based on position and signal."""
        for row_idx, (_, row) in enumerate(df.iterrows(), 2):
            for col_idx, (col_name, value) in enumerate(row.items(), 1):
                cell = worksheet.cell(row=row_idx, column=col_idx, value=value)
                
                # Apply color based on position and signal
                position = row.get('position', 'NONE')
                signal = row.get('signal', 'NONE')
                
                # Color priority: position > signal
                if position == 'LONG':
                    cell.fill = self.colors['LONG']
                elif position == 'SHORT':
                    cell.fill = self.colors['SHORT']
                elif signal == 'BUY':
                    cell.fill = self.colors['BUY']
                elif signal == 'SELL':
                    cell.fill = self.colors['SELL']
                elif signal == 'HOLD':
                    cell.fill = self.colors['HOLD']
                elif signal == 'NONE' and position == 'NONE':
                    cell.fill = self.colors['NONE']
    
    def _adjust_column_widths(self, worksheet):
        """Auto-adjust column widths for better readability."""
        for column in worksheet.columns:
            max_length = 0
            column_letter = column[0].column_letter
            
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            
            # Set width with reasonable bounds
            adjusted_width = min(max_length + 2, 20)
            worksheet.column_dimensions[column_letter].width = adjusted_width
    
    def _add_legend_worksheet(self, workbook, strategy_info: Dict[str, Any]):
        """Add legend worksheet with color explanations and strategy parameters."""
        legend_ws = workbook.create_sheet(title="Legend")
        
        # Color legend data
        legend_data = [
            ['Color Legend', ''],
            ['LONG Position', 'Light Green - Currently holding long position'],
            ['SHORT Position', 'Light Salmon - Currently holding short position'],
            ['BUY Signal', 'Lime Green - Entry signal generated'],
            ['SELL Signal', 'Light Red - Exit signal generated'],
            ['HOLD Signal', 'Gold - Holding existing position'],
            ['NONE Signal/Position', 'Light Gray - No position, no signal'],
            ['', ''],
            ['Strategy Information', ''],
        ]
        
        # Add strategy parameters if available
        if 'strategy_name' in strategy_info:
            legend_data.append(['Strategy Name', strategy_info['strategy_name']])
        
        if 'hypothesis' in strategy_info:
            legend_data.append(['Hypothesis', strategy_info['hypothesis']])
        
        if 'parameters' in strategy_info:
            legend_data.append(['', ''])
            legend_data.append(['Strategy Parameters', ''])
            for param, value in strategy_info['parameters'].items():
                legend_data.append([f'{param}', str(value)])
        
        if 'statistics' in strategy_info:
            legend_data.append(['', ''])
            legend_data.append(['Statistics', ''])
            for stat, value in strategy_info['statistics'].items():
                legend_data.append([f'{stat}', str(value)])
        
        # Add legend data to worksheet
        for row_idx, (label, description) in enumerate(legend_data, 1):
            legend_ws.cell(row=row_idx, column=1, value=label)
            legend_ws.cell(row=row_idx, column=2, value=description)
            
            # Apply colors to legend labels
            if 'LONG' in label:
                legend_ws.cell(row=row_idx, column=1).fill = self.colors['LONG']
            elif 'SHORT' in label:
                legend_ws.cell(row=row_idx, column=1).fill = self.colors['SHORT']
            elif 'BUY' in label:
                legend_ws.cell(row=row_idx, column=1).fill = self.colors['BUY']
            elif 'SELL' in label:
                legend_ws.cell(row=row_idx, column=1).fill = self.colors['SELL']
            elif 'HOLD' in label:
                legend_ws.cell(row=row_idx, column=1).fill = self.colors['HOLD']
            elif 'NONE' in label:
                legend_ws.cell(row=row_idx, column=1).fill = self.colors['NONE']
        
        # Auto-adjust legend column widths
        self._adjust_column_widths(legend_ws)
    
    def export_trade_results(
        self, 
        trades: List[Dict[str, Any]], 
        output_path: str,
        summary_stats: Optional[Dict[str, Any]] = None
    ) -> str:
        """
        Export trade results to Excel with performance analysis.
        
        Args:
            trades: List of completed trades
            output_path: Path to save Excel file
            summary_stats: Optional summary statistics
        
        Returns:
            Path to saved Excel file
        """
        if not EXCEL_AVAILABLE:
            raise RuntimeError(
                "Excel export not available. Install openpyxl: pip install openpyxl"
            )
        
        if not trades:
            raise ValueError("No trades to export")
        
        # Ensure output directory exists
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        
        # Create workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Trade Results"
        
        # Create DataFrame from trades
        df = pd.DataFrame(trades)
        
        # Add headers
        self._add_headers(ws, df.columns)
        
        # Add trade data with profit/loss coloring
        self._add_trade_data_with_colors(ws, df)
        
        # Auto-adjust column widths
        self._adjust_column_widths(ws)
        
        # Add summary statistics if provided
        if summary_stats:
            self._add_summary_worksheet(wb, summary_stats)
        
        # Save workbook
        wb.save(output_path)
        
        return output_path
    
    def _add_trade_data_with_colors(self, worksheet, df: pd.DataFrame):
        """Add trade data with profit/loss color coding."""
        for row_idx, (_, row) in enumerate(df.iterrows(), 2):
            for col_idx, (col_name, value) in enumerate(row.items(), 1):
                cell = worksheet.cell(row=row_idx, column=col_idx, value=value)
                
                # Color based on profit/loss
                pnl = row.get('pnl_usd', 0) or row.get('profit_sol', 0)
                
                if pnl > 0:
                    cell.fill = self.colors['BUY']  # Green for profit
                elif pnl < 0:
                    cell.fill = self.colors['SELL']  # Red for loss
                else:
                    cell.fill = self.colors['NONE']  # Gray for breakeven
    
    def _add_summary_worksheet(self, workbook, summary_stats: Dict[str, Any]):
        """Add summary statistics worksheet."""
        summary_ws = workbook.create_sheet(title="Summary")
        
        # Convert stats to list of tuples
        stats_data = []
        for key, value in summary_stats.items():
            # Format key for display
            display_key = key.replace('_', ' ').title()
            stats_data.append([display_key, str(value)])
        
        # Add data to worksheet
        for row_idx, (label, value) in enumerate(stats_data, 1):
            summary_ws.cell(row=row_idx, column=1, value=label)
            summary_ws.cell(row=row_idx, column=2, value=value)
        
        # Auto-adjust column widths
        self._adjust_column_widths(summary_ws)
    
    def set_custom_colors(self, color_scheme: Dict[str, str]):
        """
        Set custom color scheme.
        
        Args:
            color_scheme: Dictionary mapping signal types to hex colors
        """
        for signal_type, hex_color in color_scheme.items():
            if signal_type in self.colors:
                self.colors[signal_type] = PatternFill(
                    start_color=hex_color.replace('#', ''), 
                    end_color=hex_color.replace('#', ''), 
                    fill_type="solid"
                )


def export_strategy_to_excel(
    data: List[Dict[str, Any]], 
    output_path: str,
    strategy_info: Optional[Dict[str, Any]] = None
) -> str:
    """
    Convenience function to export strategy data to Excel.
    
    Args:
        data: Strategy signal data
        output_path: Output file path
        strategy_info: Optional strategy information
    
    Returns:
        Path to saved Excel file
    """
    exporter = ExcelExporter()
    return exporter.export_strategy_signals(data, output_path, strategy_info)


def export_trades_to_excel(
    trades: List[Dict[str, Any]], 
    output_path: str,
    summary_stats: Optional[Dict[str, Any]] = None
) -> str:
    """
    Convenience function to export trade results to Excel.
    
    Args:
        trades: List of trade results
        output_path: Output file path
        summary_stats: Optional summary statistics
    
    Returns:
        Path to saved Excel file
    """
    exporter = ExcelExporter()
    return exporter.export_trade_results(trades, output_path, summary_stats)